"""
Document generation service.
Generates DH, AE, BP (DOCX + PDF) and descriptif fournisseur.
"""
import os
import copy
import math
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Optional, Any

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from num2words import num2words

from ..config import settings


def _dh_dir(upload_dir: str) -> str:
    return upload_dir


def _format_date_fr(d: date) -> str:
    mois = ["janvier","février","mars","avril","mai","juin",
            "juillet","août","septembre","octobre","novembre","décembre"]
    return f"{d.day} {mois[d.month-1]} {d.year}"


def _montant_en_lettres(montant: float) -> str:
    try:
        entier = int(montant)
        centimes = round((montant - entier) * 100)
        text = num2words(entier, lang='fr').replace(" et ", " ").capitalize()
        if centimes > 0:
            text += f" dirhams et {num2words(centimes, lang='fr')} centimes"
        else:
            text += " dirhams"
        return text
    except Exception:
        return str(montant) + " dirhams"


def _apply_entete(doc: Document, entete_path: Optional[str], logo_path: Optional[str]):
    if entete_path and os.path.exists(entete_path):
        tpl = Document(entete_path)
        for para in tpl.paragraphs:
            if para.text.strip():
                new_para = doc.add_paragraph()
                new_para.alignment = para.alignment
                for run in para.runs:
                    new_run = new_para.add_run(run.text)
                    new_run.bold = run.bold
                    new_run.italic = run.italic
                    new_run.font.size = run.font.size
        doc.add_paragraph()
    elif logo_path and os.path.exists(logo_path):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(logo_path, width=Cm(4))
        doc.add_paragraph()


def _save_docx_as_pdf(docx_path: str) -> str:
    pdf_path = docx_path.replace(".docx", ".pdf")
    try:
        import subprocess
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(Path(docx_path).parent), docx_path],
            capture_output=True, timeout=30
        )
    except Exception:
        pass
    return pdf_path


def generate_dh(
    societe: Any,
    ao: Any,
    output_dir: str
) -> dict:
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    # Title
    title = doc.add_heading("DÉCLARATION SUR L'HONNEUR", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    ao_ref = ao.reference or "N/A"
    ao_objet = ao.objet or "N/A"
    mo_nom = (ao.maitre_ouvrage or {}).get("nom", "N/A") if ao.maitre_ouvrage else "N/A"
    societe_nom = societe.raison_sociale
    societe_rc = societe.rc or ""
    societe_ice = societe.ice or ""
    societe_gerant = societe.gerant or ""
    societe_ville = societe.ville or "Maroc"
    today = _format_date_fr(date.today())

    corps = f"""Je soussigné(e), {societe_gerant}, agissant en qualité de gérant de la société {societe_nom},
RC n° {societe_rc}, ICE n° {societe_ice}, sise à {societe.adresse or societe_ville},

déclare sur l'honneur, en réponse à l'appel d'offres N° {ao_ref} lancé par {mo_nom}
relatif à : {ao_objet}

que :

1. La société {societe_nom} est en situation régulière vis-à-vis des administrations compétentes ;

2. Elle n'est pas en état de redressement judiciaire ou de liquidation ;

3. Elle n'est pas sous le coup d'une interdiction de participer aux marchés publics ;

4. Elle est en règle vis-à-vis de l'Administration fiscale et de l'Administration des Douanes ;

5. Elle est en règle vis-à-vis des organismes de prévoyance sociale ;

6. Elle a pris connaissance du règlement de consultation et s'engage à le respecter ;

7. Les renseignements fournis dans le dossier de candidature sont exacts et sincères.

Fait à {societe_ville}, le {today}

Le Gérant,

{societe_gerant}
"""
    for line in corps.split("\n"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    docx_path = os.path.join(output_dir, "DH.docx")
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}


def generate_ae(
    societe: Any,
    ao: Any,
    montant_ht: float,
    montant_ttc: float,
    output_dir: str
) -> dict:
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    title = doc.add_heading("ACTE D'ENGAGEMENT", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    ao_ref = ao.reference or "N/A"
    ao_objet = ao.objet or "N/A"
    mo_nom = (ao.maitre_ouvrage or {}).get("nom", "N/A") if ao.maitre_ouvrage else "N/A"
    societe_nom = societe.raison_sociale
    societe_gerant = societe.gerant or ""
    societe_ville = societe.ville or "Maroc"
    today = _format_date_fr(date.today())
    montant_lettres = _montant_en_lettres(montant_ttc)
    tva = montant_ttc - montant_ht

    corps = f"""Appel d'offres N° : {ao_ref}
Objet : {ao_objet}
Maître d'ouvrage : {mo_nom}

Je soussigné(e), {societe_gerant}, gérant de la société {societe_nom},

après avoir pris connaissance du dossier d'appel d'offres et après avoir apprécié,
sous ma responsabilité, la nature et les difficultés que comprennent les prestations
à exécuter,

m'engage à réaliser lesdites prestations conformément aux conditions du Cahier des
Prescriptions Spéciales, pour les prix ci-après :

─────────────────────────────────────────────────────────
Montant HT  :  {montant_ht:,.2f} DH
TVA (20%)   :  {tva:,.2f} DH
Montant TTC :  {montant_ttc:,.2f} DH
─────────────────────────────────────────────────────────

Soit en lettres : {montant_lettres}

Le délai d'exécution est fixé selon les termes du CPS.

Fait à {societe_ville}, le {today}

Le Gérant,

{societe_gerant}
"""
    for line in corps.split("\n"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    docx_path = os.path.join(output_dir, "AE.docx")
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}


def generate_bp(
    societe: Any,
    ao: Any,
    prix_detail: list,
    montant_ht: float,
    montant_ttc: float,
    output_dir: str
) -> dict:
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    title = doc.add_heading("BORDEREAU DES PRIX — DÉTAIL ESTIMATIF", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    ao_ref = ao.reference or "N/A"
    ao_objet = ao.objet or "N/A"

    meta = doc.add_paragraph()
    meta.add_run(f"AO N° {ao_ref} — {ao_objet}").bold = True
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["N°", "Désignation", "Qté", "Unité", "Prix Unit. HT (DH)", "Montant HT (DH)"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for item in prix_detail:
        row = table.add_row().cells
        row[0].text = str(item.get("numero", ""))
        row[1].text = item.get("designation", "")
        row[2].text = str(item.get("quantite", ""))
        row[3].text = item.get("unite", "")
        row[4].text = f"{item.get('prix_unitaire', 0):,.2f}"
        row[5].text = f"{item.get('montant', 0):,.2f}"

    # Totals
    tva = montant_ttc - montant_ht
    total_row = table.add_row().cells
    total_row[0].merge(total_row[4])
    total_row[0].text = "TOTAL HT"
    total_row[0].paragraphs[0].runs[0].bold = True
    total_row[5].text = f"{montant_ht:,.2f}"

    tva_row = table.add_row().cells
    tva_row[0].merge(tva_row[4])
    tva_row[0].text = "TVA (20%)"
    tva_row[5].text = f"{tva:,.2f}"

    ttc_row = table.add_row().cells
    ttc_row[0].merge(ttc_row[4])
    ttc_row[0].text = "TOTAL TTC"
    ttc_row[0].paragraphs[0].runs[0].bold = True
    ttc_row[5].text = f"{montant_ttc:,.2f}"
    ttc_row[5].paragraphs[0].runs[0].bold = True

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    docx_path = os.path.join(output_dir, "BP.docx")
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)

    # Excel BP
    xlsx_path = _generate_bp_xlsx(prix_detail, montant_ht, montant_ttc, ao_ref, ao_objet, output_dir)

    return {"docx": docx_path, "pdf": pdf_path, "xlsx": xlsx_path}


def _generate_bp_xlsx(prix_detail: list, montant_ht: float, montant_ttc: float,
                      ao_ref: str, ao_objet: str, output_dir: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereau des Prix"

    blue = "2563EB"
    header_fill = PatternFill("solid", fgColor=blue)
    bold_font = Font(bold=True)
    white_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = f"BORDEREAU DES PRIX — AO N° {ao_ref}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = center

    ws["A2"] = ao_objet
    ws.merge_cells("A2:F2")
    ws["A2"].alignment = center
    ws["A2"].font = Font(italic=True)

    headers = ["N°", "Désignation", "Quantité", "Unité", "Prix Unitaire HT (DH)", "Montant HT (DH)"]
    widths = [6, 40, 12, 12, 22, 22]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = w

    for i, item in enumerate(prix_detail, start=5):
        data = [item.get("numero",""), item.get("designation",""),
                item.get("quantite",0), item.get("unite",""),
                item.get("prix_unitaire",0), item.get("montant",0)]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = border
            if col in (5, 6):
                cell.number_format = '#,##0.00'
            if col in (3,):
                cell.alignment = center

    last = len(prix_detail) + 5
    tva = montant_ttc - montant_ht
    for label, val in [("TOTAL HT", montant_ht), ("TVA (20%)", tva), ("TOTAL TTC", montant_ttc)]:
        ws.cell(row=last, column=1, value=label).font = bold_font
        ws.merge_cells(f"A{last}:E{last}")
        c = ws.cell(row=last, column=6, value=val)
        c.font = bold_font
        c.number_format = '#,##0.00'
        c.border = border
        last += 1

    xlsx_path = os.path.join(output_dir, "BP.xlsx")
    wb.save(xlsx_path)
    return xlsx_path


def generate_maintien_offre(
    societe: Any,
    ao: Any,
    montant_ttc: float,
    date_demande_mo: str,
    nouvelle_date_validite: str,
    output_dir: str
) -> dict:
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    doc.add_paragraph(f"{societe.ville or 'Maroc'}, le {_format_date_fr(date.today())}")
    doc.add_paragraph()

    mo = ao.maitre_ouvrage or {}
    doc.add_paragraph(f"À Monsieur {mo.get('titre', 'le Directeur')} de {mo.get('nom', '')}")
    doc.add_paragraph()

    objet = doc.add_paragraph()
    objet.add_run(f"Objet : Maintien de l'offre — AO N°{ao.reference} — {ao.objet}").bold = True
    doc.add_paragraph()

    corps = f"""Monsieur,

Suite à votre demande du {date_demande_mo}, nous avons l'honneur de vous informer
que notre société {societe.raison_sociale} maintient son offre d'un montant de
{montant_ttc:,.2f} DH TTC pour l'appel d'offres susmentionné jusqu'au {nouvelle_date_validite}.

Nous restons à votre disposition pour tout renseignement complémentaire.

Veuillez agréer, Monsieur, l'expression de nos salutations distinguées.


{societe.gerant or 'Le Gérant'}
"""
    for line in corps.split("\n"):
        doc.add_paragraph(line)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    fname = f"Maintien_Offre_{ao.reference}_{societe.code}.docx".replace("/", "-")
    docx_path = os.path.join(output_dir, fname)
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}


def generate_refus_maintien(
    societe: Any,
    ao: Any,
    output_dir: str
) -> dict:
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    doc.add_paragraph(f"{societe.ville or 'Maroc'}, le {_format_date_fr(date.today())}")
    doc.add_paragraph()

    mo = ao.maitre_ouvrage or {}
    doc.add_paragraph(f"À Monsieur {mo.get('titre', 'le Directeur')} de {mo.get('nom', '')}")
    doc.add_paragraph()

    objet = doc.add_paragraph()
    objet.add_run(f"Objet : Refus de maintien de l'offre — AO N°{ao.reference}").bold = True
    doc.add_paragraph()

    corps = f"""Monsieur,

En réponse à votre demande de prorogation du délai de validité des offres pour
l'appel d'offres N°{ao.reference} relatif à : {ao.objet},

nous vous informons que nous nous trouvons dans l'impossibilité de maintenir
notre offre au-delà de la date initiale de validité.

Nous vous remercions de l'intérêt que vous avez porté à notre société.

Veuillez agréer, Monsieur, l'expression de nos salutations distinguées.


{societe.gerant or 'Le Gérant'}
"""
    for line in corps.split("\n"):
        doc.add_paragraph(line)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    fname = f"Refus_Maintien_{ao.reference}_{societe.code}.docx".replace("/", "-")
    docx_path = os.path.join(output_dir, fname)
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}


def generate_demande_caution_definitive(
    societe: Any,
    marche: Any,
    ao: Any,
    output_dir: str
) -> dict:
    caution = math.ceil(float(marche.montant_ht or 0) * 0.03)
    caution_lettres = _montant_en_lettres(caution)

    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    doc.add_paragraph(f"{societe.ville or 'Maroc'}, le {_format_date_fr(date.today())}")
    doc.add_paragraph()
    doc.add_paragraph(f"À Monsieur {societe.titre_directeur_banque or 'Le Directeur'}")
    doc.add_paragraph(f"De {societe.banque_domiciliation or 'la Banque'}")
    doc.add_paragraph()

    objet = doc.add_paragraph()
    objet.add_run(
        f"Objet : Demande de caution définitive — Marché N°{marche.numero_marche or 'N/A'}"
    ).bold = True
    doc.add_paragraph()

    mo = ao.maitre_ouvrage or {}
    delai_exec = (ao.delai_execution or {}).get("valeur", "N/A") if ao.delai_execution else "N/A"
    delai_unit = (ao.delai_execution or {}).get("unite", "") if ao.delai_execution else ""
    delai_gar = (ao.delai_garantie or {}).get("valeur", "N/A") if ao.delai_garantie else "N/A"
    delai_gar_unit = (ao.delai_garantie or {}).get("unite", "") if ao.delai_garantie else ""

    mo_nom = mo.get("nom", "l'Administration")
    mo_nom2 = mo.get("nom", "")
    ao_objet = ao.objet or ""
    marche_num = marche.numero_marche or "N/A"
    montant_ht = float(marche.montant_ht or 0)
    rib = societe.rib or "N/A"
    gerant = societe.gerant or "Le Gérant"
    caution_fmt = f"{caution:,.0f}"
    montant_fmt = f"{montant_ht:,.2f}"

    corps = f"""Monsieur,

Nous vous demandons de bien vouloir émettre, en faveur de {mo_nom},
une caution définitive d'un montant de :

{caution_fmt} DH ({caution_lettres})

dans le cadre du marché suivant :

• Administration intéressée  : {mo_nom2}
• Objet du marché           : {ao_objet}
• Numéro du marché          : {marche_num}
• Montant HT du marché      : {montant_fmt} DH
• Délai d'exécution         : {delai_exec} {delai_unit}
• Délai de réception défin. : {delai_gar} {delai_gar_unit}
• Numéro de compte          : {rib}

Nous restons à votre disposition pour toute information complémentaire.

Veuillez agréer, Monsieur, l'expression de nos salutations distinguées.


{gerant}
"""
    for line in corps.split("\n"):
        doc.add_paragraph(line)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    fname = f"Demande_Caution_Definitive_{marche.numero_marche or 'MARCHE'}.docx".replace("/", "-")
    docx_path = os.path.join(output_dir, fname)
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path, "montant_caution": caution}


def generate_main_levee(
    societe: Any,
    marche: Any,
    date_pv_rd: str,
    output_dir: str
) -> dict:
    caution = float(marche.caution_definitive or 0)

    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    doc.add_paragraph(f"{societe.ville or 'Maroc'}, le {_format_date_fr(date.today())}")
    doc.add_paragraph()
    doc.add_paragraph(f"À Monsieur Le Directeur de {societe.banque_domiciliation or 'la Banque'}")
    doc.add_paragraph()

    objet = doc.add_paragraph()
    objet.add_run(
        f"Objet : Demande de Main Levée — Caution Définitive N° {marche.numero_marche or 'N/A'}"
    ).bold = True
    doc.add_paragraph()

    corps = f"""Monsieur,

Suite à la prononciation de la réception définitive du marché N°{marche.numero_marche or 'N/A'}
en date du {date_pv_rd}, nous vous demandons de bien vouloir procéder à la main levée
de la caution définitive d'un montant de {caution:,.0f} DH
constituée dans le cadre du marché susmentionné.

Veuillez agréer, Monsieur, l'expression de nos salutations distinguées.


{societe.gerant or 'Le Gérant'}
"""
    for line in corps.split("\n"):
        doc.add_paragraph(line)

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    fname = f"Main_Levee_Caution_{marche.numero_marche or 'MARCHE'}.docx".replace("/", "-")
    docx_path = os.path.join(output_dir, fname)
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}


def generate_excel_enregistrement(
    marche: Any,
    ao: Any,
    periodicite: str,
    output_dir: str
) -> str:
    import math as m
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enregistrement Marché"

    blue = "2563EB"
    header_fill = PatternFill("solid", fgColor=blue)
    white_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 40

    ws["A1"] = "ENREGISTREMENT COMPTABLE DU MARCHÉ"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 30

    rows = [
        ("Référence du marché", marche.numero_marche or ""),
        ("Objet", ao.objet or ""),
        ("Maître d'ouvrage", (ao.maitre_ouvrage or {}).get("nom", "") if ao.maitre_ouvrage else ""),
        ("Montant HT du marché (DH)", float(marche.montant_ht or 0)),
        ("Montant TTC du marché (DH)", float(marche.montant_ttc or 0)),
        ("Nature de la prestation", ao.domaine or ""),
        ("Date d'approbation", ""),
        ("Date d'exécution (OS)", ""),
        ("Date du premier paiement", ""),
        ("Date du dernier paiement", ""),
        ("Périodicité du paiement", periodicite),
        ("Caution définitive (3% HT)", m.ceil(float(marche.montant_ht or 0) * 0.03)),
        ("Retenue de garantie (7% HT)", round(float(marche.montant_ht or 0) * 0.07, 2)),
    ]

    for i, (label, val) in enumerate(rows, start=3):
        a = ws.cell(row=i, column=1, value=label)
        a.font = bold
        a.border = border
        a.fill = PatternFill("solid", fgColor="F3F4F6")
        b = ws.cell(row=i, column=2, value=val)
        b.border = border
        if isinstance(val, float):
            b.number_format = '#,##0.00'

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    fname = f"Enregistrement_{marche.numero_marche or 'MARCHE'}_{date.today().strftime('%Y%m%d')}.xlsx".replace("/", "-")
    path = os.path.join(output_dir, fname)
    wb.save(path)
    return path


def generate_descriptif_fournisseur(
    societe: Any,
    ao: Any,
    articles: list,
    langue: str,
    output_dir: str
) -> dict:
    from . import ai_service
    doc = Document()
    _apply_entete(doc, societe.entete_path, societe.logo_path)

    title = doc.add_heading(
        "DEMANDE DE PRIX — DESCRIPTIF TECHNIQUE" if langue == "fr"
        else "REQUEST FOR QUOTATION — TECHNICAL SPECIFICATIONS",
        level=1
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_str = _format_date_fr(date.today()) if langue == "fr" else date.today().strftime("%B %d, %Y")
    doc.add_paragraph(f"{'Date :' if langue == 'fr' else 'Date:'} {date_str}")
    doc.add_paragraph()

    note = doc.add_paragraph()
    if langue == "fr":
        note.add_run("Merci de nous communiquer vos meilleurs prix pour les articles suivants :").italic = True
    else:
        note.add_run("Please provide your best prices for the following items:").italic = True
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    cols_fr = ["N°", "Désignation", "Spécifications Techniques", "Quantité", "Unité"]
    cols_en = ["#", "Description", "Technical Specifications", "Quantity", "Unit"]
    headers = cols_fr if langue == "fr" else cols_en
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True

    for article in articles:
        specs = article.get("specifications_techniques", "")
        if langue == "en" and specs:
            try:
                specs = ai_service.translate_specs_to_english(specs)
            except Exception:
                pass
        row = table.add_row().cells
        row[0].text = str(article.get("numero", ""))
        row[1].text = article.get("designation", "")
        row[2].text = specs
        row[3].text = str(article.get("quantite", ""))
        row[4].text = article.get("unite", "")

    Path(output_dir).mkdir(parents=True, exist_ok=True)
    article_label = "TOUS" if len(articles) > 1 else articles[0].get("designation", "ART")[:20]
    fname = f"Descriptif_Technique_{article_label}_{date.today().strftime('%Y%m%d')}.docx".replace(" ", "_").replace("/", "-")
    docx_path = os.path.join(output_dir, fname)
    doc.save(docx_path)
    pdf_path = _save_docx_as_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path}
