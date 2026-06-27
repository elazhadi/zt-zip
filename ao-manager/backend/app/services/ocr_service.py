import os
import io
from pathlib import Path
from typing import Optional
import pdfplumber
from pypdf import PdfReader
from PIL import Image

try:
    import pytesseract
    TESSERACT_AVAILABLE = True
except ImportError:
    TESSERACT_AVAILABLE = False


def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        if text_parts:
            return "\n\n".join(text_parts)
    except Exception:
        pass

    # Fallback: pypdf
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
        if text_parts:
            return "\n\n".join(text_parts)
    except Exception:
        pass

    # Fallback: OCR
    if TESSERACT_AVAILABLE:
        return ocr_pdf(file_path)
    return ""


def ocr_pdf(file_path: str) -> str:
    import subprocess
    text_parts = []
    try:
        result = subprocess.run(
            ["pdfimages", "-png", file_path, "/tmp/ao_page"],
            capture_output=True, timeout=60
        )
        import glob
        images = sorted(glob.glob("/tmp/ao_page*.png"))
        for img_path in images:
            try:
                img = Image.open(img_path)
                text = pytesseract.image_to_string(img, lang="fra+ara", config="--psm 6")
                text_parts.append(text)
                os.unlink(img_path)
            except Exception:
                pass
    except Exception:
        pass
    return "\n\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts)


def extract_text_from_image(file_path: str) -> str:
    if not TESSERACT_AVAILABLE:
        return ""
    try:
        img = Image.open(file_path)
        return pytesseract.image_to_string(img, lang="fra+ara", config="--psm 6")
    except Exception:
        return ""


def extract_text_from_xlsx(file_path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"=== Feuille: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(c) if c is not None else "" for c in row)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts)


def extract_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    elif ext in (".jpg", ".jpeg", ".png", ".tiff", ".bmp"):
        return extract_text_from_image(file_path)
    elif ext in (".xlsx", ".xls"):
        return extract_text_from_xlsx(file_path)
    return ""
